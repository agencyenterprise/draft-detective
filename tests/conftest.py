"""Shared test utilities for all agent tests.

This module provides reusable utilities that work across all agent test suites:
- Path resolution
- Document loading
- A small DOCX with Word equations
"""

import os
import shutil
from pathlib import Path

from docx import Document
from lxml import etree  # type: ignore[import-untyped]
from xxhash import xxh128

from lib.config.env import config
from lib.services.files.document import create_file_document_from_path

TESTS_DIR = Path(__file__).parent

_OMML_NS = (
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
)
# x subscript i: MarkItDown converts this to LaTeX.
_OMML_SUBSCRIPT = (
    f"<m:oMath {_OMML_NS}><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e>"
    "<m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:oMath>"
)
# ln(N): MarkItDown 0.1.5 raises "Not support func ln".
_OMML_NATURAL_LOG = (
    f"<m:oMath {_OMML_NS}><m:func><m:fName><m:r><m:t>ln</m:t></m:r></m:fName>"
    "<m:e><m:r><m:t>N</m:t></m:r></m:e></m:func></m:oMath>"
)


def data_path(path: str) -> str:
    """
    Convert relative test data path to absolute path.

    Args:
        path: Relative path from tests/ directory (e.g., "data/common_knowledge/main.md")

    Returns:
        Absolute path to the file
    """
    return str(TESTS_DIR / path)


async def create_test_file_document_from_path(path: str):
    """
    Load a single document from test data.

    Copies the test file to the uploads directory with an xxhash-based filename,
    similar to how uploaded files are handled in production.

    Args:
        path: Relative path from tests/ directory

    Returns:
        FileDocument object with markdown content
    """
    source_path = data_path(path)

    with open(source_path, "rb") as f:
        content = f.read()

    xxhash = xxh128(content).hexdigest()

    filename = os.path.basename(source_path)
    file_extension = os.path.splitext(filename)[1]

    upload_dir = config.FILE_UPLOADS_MOUNT_PATH
    dest_path = os.path.join(upload_dir, xxhash + file_extension)

    # Copy file to uploads directory atomically (safe for parallel test execution)
    # Use a temp file + rename pattern to avoid race conditions
    os.makedirs(upload_dir, exist_ok=True)
    if not os.path.exists(dest_path):
        temp_path = dest_path + f".{os.getpid()}.tmp"
        try:
            shutil.copy2(source_path, temp_path)
            os.rename(temp_path, dest_path)  # Atomic on POSIX
        except FileExistsError:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        except OSError:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    original_file_name = filename

    return await create_file_document_from_path(
        dest_path,
        file_id=path,
        original_file_name=original_file_name,
        markdown_convert=True,
    )


def write_docx_with_equations(path: Path) -> str:
    """Write a DOCX with a convertible inline equation, an ``ln`` one that
    MarkItDown cannot convert, and a display equation. Returns its path."""
    doc = Document()
    for before, omath, after in (
        ("where ", _OMML_SUBSCRIPT, " is the score for item i."),
        ("The entropy uses ", _OMML_NATURAL_LOG, " for the base."),
    ):
        paragraph = doc.add_paragraph(before)
        paragraph._element.append(etree.fromstring(omath))
        paragraph.add_run(after)
    display = f"<m:oMathPara {_OMML_NS}>{_OMML_SUBSCRIPT}</m:oMathPara>"
    doc.add_paragraph()._element.append(etree.fromstring(display))
    doc.add_paragraph("After the display equation.")
    doc.save(str(path))
    return str(path)
