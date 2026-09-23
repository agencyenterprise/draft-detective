"""Offline tests: source fidelity, independent calls and deterministic accounting."""

import asyncio
import json

import httpx
import pytest
from docx import Document
from langchain_core.messages import AIMessage
from langchain_openai import ChatOpenAI

from lib.agents.abbreviation_checker_structured import (
    ChunkExtraction,
    ChunkOccurrence,
    GlossaryEntry,
    StructuredAbbreviationChecker,
    chunk_text,
    reconcile_chunks,
)


def extraction(**kwargs):
    return ChunkExtraction(
        **{
            "occurrences": [],
            "glossary": [],
            "glossary_unit_ids": [],
            "abbreviations_section_found": False,
            **kwargs,
        }
    )


def occurrence(unit_id=0, surface="AI", match_number=1, **kwargs):
    return ChunkOccurrence(
        unit_id=unit_id,
        surface=surface,
        match_number=match_number,
        inline_definition=kwargs.get("inline_definition", ""),
        ignored_reason=kwargs.get("ignored_reason", ""),
    )


@pytest.mark.asyncio
async def test_docx_uses_existing_markdown_converter_including_tables(tmp_path):
    from lib.services.converters.base import convert_to_markdown

    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Artificial intelligence (AI).")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "AI"
    table.cell(0, 1).text = "artificial intelligence"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "LLM"
    nested = merged.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested"
    nested.cell(0, 1).text = "text"
    doc.add_paragraph("After table")
    path = tmp_path / "document.docx"
    doc.save(path)

    result = await StructuredAbbreviationChecker(FakeModel()).check_docx(path)
    text = result.source_text
    assert text == await convert_to_markdown(str(path), converter="markitdown")
    assert "# Introduction" in text
    assert "AI" in text and "artificial intelligence" in text
    assert "|" in text  # Tables are Markdown tables, not a separate XML projection.
    assert "Nested" in text
    assert text.index("Nested") < text.index("After table")


def test_chunks_have_disjoint_ownership_bounded_context_and_source_offsets():
    text = "# References\n" + "AI item " * 100 + "\n\nEnd"
    chunks = chunk_text(text, chunk_chars=128, context_chars=64)
    units = [unit for chunk in chunks for unit in chunk.core]
    assert len(units) == len({unit.id for unit in units})
    for chunk in chunks:
        assert sum(len(u.text) for u in chunk.core) <= 128
        assert sum(len(u.text) for u in chunk.before) <= 64
        assert sum(len(u.text) for u in chunk.after) <= 64
    for unit in units:
        assert (
            text.splitlines()[unit.line - 1][unit.offset : unit.offset + len(unit.text)]
            == unit.text
        )
        assert unit.heading == "# References"
    assert units[-1].line == 4
    assert "".join(u.text for u in units if u.line == 2) == text.splitlines()[1]


def test_reconcile_repeats_plural_definitions_global_numbering_and_glossary():
    text = "large language models (LLMs), LLM and LLM.\n# Abbreviations\nLLM\tlarge language model"
    chunks = chunk_text(text)
    first = occurrence(surface="LLMs", inline_definition="large language models")
    result = reconcile_chunks(
        text,
        chunks,
        [
            extraction(
                occurrences=[
                    occurrence(surface="LLM", match_number=2),
                    first,
                    occurrence(surface="LLM"),
                    first,  # Identical duplicate reporting is not a second occurrence.
                    occurrence(
                        unit_id=2, surface="LLM"
                    ),  # Glossary catalogue excluded.
                ],
                glossary=[
                    GlossaryEntry(
                        unit_id=2, surface="LLM", definition="large language model"
                    )
                ],
                glossary_unit_ids=[1, 2],
                abbreviations_section_found=True,
            )
        ],
    )
    assert [item.abbr for item in result.abbreviations] == ["LLM"] * 3
    assert [item.occurrence_number for item in result.abbreviations] == [1, 2, 3]
    assert [item.inline_definition for item in result.abbreviations] == [
        "large language models",
        "",
        "",
    ]
    assert all(
        item.abbreviations_section_definition == "large language model"
        for item in result.abbreviations
    )
    assert result.output.abbreviations_section_found
    assert result.as_agent_result()[1] == result.abbreviations


def test_glossary_is_propagated_across_chunks():
    text = "AI\n" + "padding " * 40 + "\n# Acronyms\nAI artificial intelligence"
    chunks = chunk_text(text, chunk_chars=128, context_chars=64)
    responses = []
    for chunk in chunks:
        mentions, glossary, glossary_units = [], [], []
        for unit in chunk.core:
            if unit.line == 1:
                mentions.append(occurrence(unit_id=unit.id))
            if unit.line == 4:
                glossary.append(
                    GlossaryEntry(
                        unit_id=unit.id,
                        surface="AI",
                        definition="artificial intelligence",
                    )
                )
                glossary_units.append(unit.id)
        responses.append(
            extraction(
                occurrences=mentions,
                glossary=glossary,
                glossary_unit_ids=glossary_units,
            )
        )
    result = reconcile_chunks(text, chunks, responses)
    assert len(chunks) > 1
    assert (
        result.abbreviations[0].abbreviations_section_definition
        == "artificial intelligence"
    )


def test_conflicting_glossary_definitions_not_silently_chosen():
    text = "AI\n# Glossary\nAI artificial intelligence\nAI alternate interpretation"
    result = reconcile_chunks(
        text,
        chunk_text(text),
        [
            extraction(
                occurrences=[occurrence()],
                glossary_unit_ids=[1, 2, 3],
                glossary=[
                    GlossaryEntry(
                        unit_id=2, surface="AI", definition="artificial intelligence"
                    ),
                    GlossaryEntry(
                        unit_id=3, surface="AI", definition="alternate interpretation"
                    ),
                ],
            )
        ],
    )
    assert result.abbreviations[0].abbreviations_section_definition is None
    assert "Conflicting glossary definitions for AI" in result.warnings[0]


@pytest.mark.parametrize(
    "bad",
    [
        occurrence(surface="invented"),
        occurrence(match_number=2),
        occurrence(match_number=0),
        occurrence(unit_id=999),
        occurrence(inline_definition="imaginary definition"),
    ],
)
def test_ungrounded_occurrences_are_rejected(bad):
    with pytest.raises(ValueError):
        reconcile_chunks(
            "AI",
            chunk_text("AI"),
            [extraction(occurrences=[bad])],
            strict_grounding=True,
        )


def test_conflicting_duplicate_reports_are_rejected():
    with pytest.raises(ValueError, match="Conflicting reports"):
        reconcile_chunks(
            "AI",
            chunk_text("AI"),
            [
                extraction(
                    occurrences=[
                        occurrence(),
                        occurrence(ignored_reason="Heading"),
                    ]
                )
            ],
            strict_grounding=True,
        )


def test_nonexistent_third_match_does_not_abort_two_grounded_occurrences():
    result = reconcile_chunks(
        "AI and AI",
        chunk_text("AI and AI"),
        [extraction(occurrences=[occurrence(match_number=i) for i in (1, 2, 3)])],
    )
    assert len(result.abbreviations) == 2
    assert [item.occurrence_number for item in result.abbreviations] == [1, 2]
    assert len(result.warnings) == 1
    assert "reported 3, source has 2" in result.warnings[0]


def test_bad_records_are_diagnosed_without_losing_valid_mentions():
    result = reconcile_chunks(
        "AI LLM",
        chunk_text("AI LLM"),
        [
            extraction(
                occurrences=[
                    occurrence(surface="invented"),
                    occurrence(unit_id=999),
                    occurrence(inline_definition="invented expanded name"),
                    occurrence(surface="LLM"),
                ],
                glossary_unit_ids=[999],
                glossary=[
                    GlossaryEntry(unit_id=999, surface="AI", definition="invented")
                ],
            )
        ],
    )
    assert [item.abbr for item in result.abbreviations] == ["AI", "LLM"]
    assert all(item.inline_definition == "" for item in result.abbreviations)
    assert len(result.warnings) == 5


def test_conflicting_annotations_cleared_without_duplicating_occurrence():
    result = reconcile_chunks(
        "AI",
        chunk_text("AI"),
        [extraction(occurrences=[occurrence(), occurrence(ignored_reason="Heading")])],
    )
    assert len(result.abbreviations) == 1
    assert not result.abbreviations[0].ignored
    assert "Conflicting reports" in result.warnings[0]


def test_case_boundaries_and_ignored_mentions():
    text = "# AI\nAIDS AI AIs paid ai"
    result = reconcile_chunks(
        text,
        chunk_text(text),
        [
            extraction(
                occurrences=[
                    occurrence(ignored_reason="Heading"),
                    occurrence(unit_id=1, surface="AIDS"),
                    occurrence(unit_id=1),
                    occurrence(unit_id=1, surface="AIs"),
                ]
            )
        ],
    )
    assert [item.abbr for item in result.abbreviations] == ["AI", "AIDS", "AI", "AI"]
    assert result.abbreviations[0].ignored
    assert result.abbreviations[-1].occurrence_number == 3


@pytest.mark.parametrize(
    "kwargs",
    [
        {"chunk_chars": 0},
        {"context_chars": -1},
    ],
)
def test_invalid_budgets(kwargs):
    with pytest.raises(ValueError):
        chunk_text("AI", **kwargs)


def test_long_unbroken_token_fails_without_truncation():
    with pytest.raises(ValueError, match="Unbroken token"):
        chunk_text("x" * 200, chunk_chars=128, context_chars=0)


def test_split_line_mentions_have_distinct_source_anchors_and_global_ordinals():
    text = "AI " * 100
    chunks = chunk_text(text, chunk_chars=128, context_chars=64)
    responses = [
        extraction(
            occurrences=[
                occurrence(unit_id=unit.id, match_number=index + 1)
                for unit in chunk.core
                for index in range(unit.text.count("AI"))
            ]
        )
        for chunk in chunks
    ]
    result = reconcile_chunks(text, chunks, responses)
    assert len(result.abbreviations) == 100
    assert [item.occurrence_number for item in result.abbreviations] == list(
        range(1, 101)
    )
    assert all(item.line_start == item.line_end == 1 for item in result.abbreviations)


def test_inline_definition_can_come_from_neighboring_context():
    text = "padding " * 13 + "\nartificial intelligence\n(AI)"
    chunks = chunk_text(text, chunk_chars=128, context_chars=64)
    last = chunks[-1]
    assert last.core[0].text == "(AI)"
    assert any("artificial intelligence" in unit.text for unit in last.before)
    responses = [extraction() for _ in chunks]
    responses[-1] = extraction(
        occurrences=[
            occurrence(
                unit_id=last.core[0].id, inline_definition="artificial intelligence"
            )
        ]
    )
    result = reconcile_chunks(text, chunks, responses)
    assert result.abbreviations[0].inline_definition == "artificial intelligence"


class FakeModel:
    cache = False

    def __init__(self, fail=False):
        self.fail = fail
        self.calls = []
        self.active = 0
        self.peak = 0
        self.cancelled = 0

    def with_structured_output(self, schema, **kwargs):
        assert schema is ChunkExtraction
        assert kwargs == {"method": "json_schema", "strict": True, "include_raw": True}
        return self

    async def ainvoke(self, prompt, config=None):
        units = json.loads(prompt[1].content)
        self.calls.append((units, config))
        self.active += 1
        self.peak = max(self.peak, self.active)
        try:
            if self.fail and len(self.calls) == 1:
                await asyncio.sleep(0)
                return {
                    "parsed": None,
                    "parsing_error": ValueError("truncated"),
                    "raw": AIMessage(content=""),
                }
            await asyncio.sleep(0.01)
            return {
                "parsed": extraction(),
                "parsing_error": None,
                "raw": AIMessage(content="{}"),
            }
        except asyncio.CancelledError:
            self.cancelled += 1
            raise
        finally:
            self.active -= 1


@pytest.mark.asyncio
async def test_one_call_per_chunk_bounded_concurrency_and_fresh_calls_each_run():
    model = FakeModel()
    checker = StructuredAbbreviationChecker(
        model, chunk_chars=128, context_chars=64, max_concurrency=2
    )
    text = "AI document text\n" * 40
    first = await checker.check_text(text, config={"tags": ["epoch-1"]})
    second = await checker.check_text(text)
    assert first.chunk_count > 2
    assert len(model.calls) == 2 * first.chunk_count
    assert first.abbreviations == second.abbreviations == []
    assert model.peak == 2
    assert model.calls[0][1] == {"tags": ["epoch-1"]}
    assert len(first.messages) == first.chunk_count * 3


@pytest.mark.asyncio
async def test_failure_cancels_other_calls_without_returning_partial_results():
    model = FakeModel(fail=True)
    checker = StructuredAbbreviationChecker(
        model, chunk_chars=128, context_chars=64, max_concurrency=2
    )
    with pytest.raises(ExceptionGroup, match="TaskGroup") as exc:
        await checker.check_text("AI document text\n" * 40)
    assert "chunk 0" in str(exc.value.exceptions[0])
    assert model.active == 0
    assert model.cancelled > 0


@pytest.mark.asyncio
async def test_empty_document_does_not_call_model():
    model = FakeModel()
    result = await StructuredAbbreviationChecker(model).check_text("\n  \n")
    assert result.chunk_count == 0
    assert result.abbreviations == []
    assert not model.calls


@pytest.mark.asyncio
async def test_real_model_native_schema_request_and_cache_disabled_without_network():
    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "id": "fake-completion",
                "object": "chat.completion",
                "created": 0,
                "model": "gpt-5.6-terra",
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {
                            "role": "assistant",
                            "content": extraction(
                                occurrences=[occurrence()]
                            ).model_dump_json(),
                        },
                    }
                ],
                "usage": {
                    "prompt_tokens": 10,
                    "completion_tokens": 10,
                    "total_tokens": 20,
                },
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        model = ChatOpenAI(
            model="gpt-5.6-terra",
            api_key="test-placeholder",
            max_retries=0,
            http_async_client=client,
            cache=True,
        )
        checker = StructuredAbbreviationChecker(model)
        for _ in range(2):
            result = await checker.check_text("AI")
            assert result.abbreviations[0].abbr == "AI"
    assert model.cache is True  # The caller's model was not changed.
    assert len(requests) == 2
    for request in requests:
        assert "tools" not in request
        assert request["response_format"]["type"] == "json_schema"
        assert request["response_format"]["json_schema"]["strict"] is True


def test_unsupported_native_output_does_not_silently_fall_back_to_tools():
    model = ChatOpenAI(model="gpt-4", api_key="test-placeholder")
    with pytest.raises(ValueError, match="native JSON-schema"):
        StructuredAbbreviationChecker(model)
