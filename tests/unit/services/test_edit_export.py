"""Tests for the proposed-edit half of a DOCX export, driven from Issue rows."""

import shutil
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Sequence

import pytest
from docx import Document as PythonDocxDocument
from docx_editor import BatchOperationError, Document as EditorDocument

from lib.models.issue import Issue
from lib.models.issue_edit import IssueEdit, IssueEditStatus
from lib.services.docx import edit_export
from lib.services.docx.edit_export import (
    apply_edit_export,
    describe_edit_export,
    plan_edit_export,
)
from lib.workflows.models import SeverityEnum, WorkflowRunType

_PARAGRAPH = "The chapter reports a 14% rise in output for 2019."
_MARKDOWN = _PARAGRAPH


@pytest.fixture
def docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph(_PARAGRAPH)
    path = tmp_path / "main.docx"
    document.save(str(path))
    return path


def _edit(
    original_text: str,
    replacement_text: str,
    *,
    status: IssueEditStatus = IssueEditStatus.PROPOSED,
    start_line: int = 1,
) -> IssueEdit:
    return IssueEdit(
        id=uuid.uuid4(),
        issue_id=uuid.uuid4(),
        position=0,
        original_text=original_text,
        replacement_text=replacement_text,
        start_line=start_line,
        end_line=start_line,
        rationale="the figure is wrong",
        status=status,
    )


def _issue(edits: Sequence[IssueEdit]) -> Issue:
    now = datetime.now(UTC)
    issue = Issue(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        workflow_run_id=uuid.uuid4(),
        issue_hash=uuid.uuid4().hex[:64],
        title="Reference has incorrect fields",
        description="The figure does not match the source.",
        severity=SeverityEnum.HIGH,
        workflow_type=WorkflowRunType.REFERENCE_VALIDATION_V2,
        start_line=1,
        end_line=1,
        created_at=now,
        updated_at=now,
    )
    issue.edit_rows = list(edits)
    return issue


def _visible(path: Path) -> str:
    doc = EditorDocument.open(
        path, author="Reader", workspace_dir=str(path.parent / "read-ws")
    )
    try:
        return doc.get_visible_text()
    finally:
        doc.close()


def _revisions(path: Path) -> list[tuple[str, str]]:
    doc = EditorDocument.open(
        path, author="Reader", workspace_dir=str(path.parent / "revision-ws")
    )
    try:
        return [(r.type, r.text) for r in doc.list_revisions()]
    finally:
        doc.close()


@pytest.mark.asyncio
async def test_an_issues_edit_is_planned_noted_and_written(
    docx_path: Path, tmp_path: Path
):
    issue = _issue([_edit("14% rise", "18% rise")])

    plan = await plan_edit_export(
        [issue], _MARKDOWN, str(docx_path), {0: (1, 1)}, workspace_root=str(tmp_path)
    )
    await apply_edit_export(
        str(docx_path), plan, "project-1", workspace_root=str(tmp_path)
    )

    assert plan.notes_for(issue.id) == [
        'Proposed edit: "14% rise" → "18% rise"\n'
        "the figure is wrong\n"
        "Applied below as a tracked change."
    ]
    assert "reports a 18% rise" in _visible(docx_path)


@pytest.mark.asyncio
async def test_without_a_paragraph_mapping_nothing_is_written_and_the_note_says_so(
    docx_path: Path, tmp_path: Path
):
    issue = _issue([_edit("14% rise", "18% rise")])
    before = docx_path.read_bytes()

    plan = await plan_edit_export(
        [issue], _MARKDOWN, str(docx_path), {}, workspace_root=str(tmp_path)
    )
    await apply_edit_export(
        str(docx_path), plan, "project-1", workspace_root=str(tmp_path)
    )

    assert plan.planned == []
    assert plan.notes_for(issue.id)[0].endswith(
        "Not applied as a tracked change: the quoted text could not be "
        "matched in this paragraph."
    )
    assert docx_path.read_bytes() == before


@pytest.mark.asyncio
async def test_a_rejected_edit_is_not_mentioned_at_all(docx_path: Path, tmp_path: Path):
    issue = _issue([_edit("14% rise", "18% rise", status=IssueEditStatus.REJECTED)])

    plan = await plan_edit_export(
        [issue], _MARKDOWN, str(docx_path), {0: (1, 1)}, workspace_root=str(tmp_path)
    )

    assert plan.planned == []
    assert plan.notes_for(issue.id) == []


@pytest.mark.asyncio
async def test_an_issue_without_edits_plans_nothing(docx_path: Path, tmp_path: Path):
    issue = _issue([])

    plan = await plan_edit_export(
        [issue], _MARKDOWN, str(docx_path), {0: (1, 1)}, workspace_root=str(tmp_path)
    )

    assert plan.planned == []
    assert plan.outcomes == {}
    assert plan.notes_by_issue == {}


_TWO_PARAGRAPHS = [
    "The chapter reports a 14% rise in output for 2019.",
    "The appendix lists every source consulted in full.",
]
_TWO_PARAGRAPH_MARKDOWN = "\n".join(_TWO_PARAGRAPHS)
_TWO_PARAGRAPH_RANGES = {0: (1, 1), 1: (2, 2)}


@pytest.fixture
def two_paragraph_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    for text in _TWO_PARAGRAPHS:
        document.add_paragraph(text)
    path = tmp_path / "two-paragraphs.docx"
    document.save(str(path))
    return path


@pytest.fixture
def reject_the_appendix(monkeypatch: pytest.MonkeyPatch) -> None:
    """Make docx-editor refuse the appendix paragraph's batch, and only that one.

    Stands in for anything the library rejects at apply time -- a stale hash, a
    search string the paragraph no longer carries. It bites during the
    rehearsal; the real run never asks for the operation again.
    """
    original = EditorDocument.batch_edit

    def batch_edit(self, operations, **kwargs):  # type: ignore[no-untyped-def]
        if any(operation.find == "every source" for operation in operations):
            raise BatchOperationError(0, "text not found in paragraph")
        return original(self, operations, **kwargs)

    monkeypatch.setattr(EditorDocument, "batch_edit", batch_edit)


class TestTheRehearsalDecidesWhatTheCommentsClaim:
    async def _plan(self, path: Path, issues, workspace: Path):
        return await plan_edit_export(
            issues,
            _TWO_PARAGRAPH_MARKDOWN,
            str(path),
            _TWO_PARAGRAPH_RANGES,
            workspace_root=str(workspace),
        )

    @pytest.mark.asyncio
    async def test_an_edit_the_library_refuses_is_noted_as_failed_and_dropped(
        self, two_paragraph_docx_path: Path, tmp_path: Path, reject_the_appendix: None
    ):
        writable = _issue([_edit("14% rise", "18% rise")])
        refused = _issue(
            [_edit("every source", "every source and dataset", start_line=2)]
        )

        plan = await self._plan(two_paragraph_docx_path, [writable, refused], tmp_path)
        await apply_edit_export(
            str(two_paragraph_docx_path),
            plan,
            "project-1",
            workspace_root=str(tmp_path),
        )

        (refused_note,) = plan.notes_for(refused.id)
        assert refused_note.endswith(
            "Not applied as a tracked change: text not found in paragraph."
        )
        assert [edit.edit_id for edit in plan.planned] == [writable.edit_rows[0].id]
        assert plan.notes_for(writable.id) == [
            'Proposed edit: "14% rise" → "18% rise"\n'
            "the figure is wrong\n"
            "Applied below as a tracked change."
        ]
        visible = _visible(two_paragraph_docx_path)
        assert "reports a 18% rise" in visible
        assert "lists every source consulted" in visible

    @pytest.mark.asyncio
    async def test_the_rehearsal_leaves_the_original_alone_and_is_not_applied_twice(
        self, two_paragraph_docx_path: Path, tmp_path: Path
    ):
        issue = _issue([_edit("14% rise", "18% rise")])
        before = two_paragraph_docx_path.read_bytes()

        plan = await self._plan(two_paragraph_docx_path, [issue], tmp_path)

        assert two_paragraph_docx_path.read_bytes() == before

        await apply_edit_export(
            str(two_paragraph_docx_path),
            plan,
            "project-1",
            workspace_root=str(tmp_path),
        )

        assert [kind for kind, _ in _revisions(two_paragraph_docx_path)] == [
            "deletion",
            "insertion",
        ]
        assert _visible(two_paragraph_docx_path).count("18% rise") == 1


class TestTheCommentsOnlyMode:
    def test_every_edit_is_described_without_opening_the_document(
        self, monkeypatch: pytest.MonkeyPatch
    ):
        def refuse(*args: object, **kwargs: object) -> None:
            raise AssertionError("the comments-only mode must not open the document")

        monkeypatch.setattr(edit_export, "plan_tracked_changes", refuse)
        monkeypatch.setattr(edit_export, "apply_tracked_changes", refuse)
        issue = _issue([_edit("14% rise", "18% rise")])

        export = describe_edit_export([issue])

        assert export.planned == []
        assert export.notes_for(issue.id) == [
            'Proposed edit: "14% rise" → "18% rise"\n'
            "the figure is wrong\n"
            "Not applied as a tracked change: tracked changes were not "
            "requested for this export."
        ]

    def test_a_rejected_edit_is_still_left_out(self):
        issue = _issue(
            [
                _edit("14% rise", "18% rise"),
                _edit("for 2019", "for 2021", status=IssueEditStatus.REJECTED),
            ]
        )

        export = describe_edit_export([issue])

        assert len(export.outcomes) == 1
        assert len(export.notes_for(issue.id)) == 1
        assert export.outcomes[issue.edit_rows[0].id].status == "skipped"

    def test_an_issue_without_edits_says_nothing(self):
        issue = _issue([])

        export = describe_edit_export([issue])

        assert export.outcomes == {}
        assert export.notes_by_issue == {}


class TestAnEmptyPlanCostsNothing:
    @pytest.mark.asyncio
    async def test_nothing_to_write_means_no_copy_and_no_rehearsal(
        self, docx_path: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        def refuse(*args: object, **kwargs: object) -> None:
            raise AssertionError("an empty plan must not be rehearsed")

        monkeypatch.setattr(shutil, "copyfile", refuse)
        monkeypatch.setattr(edit_export, "apply_tracked_changes", refuse)
        # The quote is nowhere on its line, so the pre-flight plans nothing.
        issue = _issue([_edit("a 40% drop", "a 30% drop")])

        export = await plan_edit_export(
            [issue],
            _MARKDOWN,
            str(docx_path),
            {0: (1, 1)},
            workspace_root=str(tmp_path),
        )

        assert export.planned == []
        assert export.notes_for(issue.id)[0].endswith(
            "Not applied as a tracked change: the quoted text could not be "
            "matched in this paragraph."
        )
