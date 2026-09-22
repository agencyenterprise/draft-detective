"""Shared fixtures and helpers for the tracked-changes export tests.

A plain module rather than a conftest, imported by name: the suites are split
by what they check -- planning, paragraph membership, replacements, conflicts,
applying, note wording -- and every one of them builds the same kind of
document, the same edit rows and the same plan-then-apply call.

Everything here runs against a real .docx built in ``tmp_path`` and read back
with docx-editor, so the redlines are checked as Word would see them. A test
module that wants the `docx_path` fixture imports the name into its own
namespace, where pytest finds it.
"""

import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import pytest
from docx import Document as PythonDocxDocument
from docx_editor import Document as EditorDocument

from lib.models.issue_edit import IssueEdit, IssueEditStatus
from lib.services.docx.tracked_changes import EditOutcome, plan_tracked_changes
from lib.services.edit_conflicts import EditCandidate, resolve_edit_conflicts
from lib.services.markdown_text import (
    locate_in_paragraph,
    render_replacement_text,
    rendered_span,
)
from lib.workflows.models import SeverityEnum
from lib.workflows.simple_deep_agent.edit_anchoring import (
    document_lines,
    normalize_whitespace,
)

# The document under test. A non-breaking space inside prose and a phrase
# repeated in one paragraph are both ordinary in converted DOCX documents; the
# table comes first on purpose, so a naive index alignment between python-docx
# and docx-editor would redline a table cell.
BODY = [
    "The Energy\u00a0Supply chapter reports a 14% rise in output for 2019.",
    "Figure 3 and Figure 3 close the section.",
    "Shared caption text",
]
CELLS = ["Shared caption text", "Other cell"]

# The markdown the edits were anchored against. Line 1 carries a trailing
# clause the DOCX paragraph does not, standing in for conversion drift.
MARKDOWN = "\n".join(
    [
        "The Energy\u00a0Supply chapter reports a 14% rise in output for 2019, per Table 2.",
        "Figure 3 and **Figure 3** close the section.",
        "Shared caption text",
    ]
)
PARAGRAPH_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 1), 1: (2, 2), 2: (3, 3)}


# Why an edit was turned down: the line and the mapped paragraph are not the
# same passage. Reported by more than one of the suites.
OTHER_PASSAGE_DETAIL = (
    "the edit's line is not part of the mapped paragraph (a table or nested block)"
)


@pytest.fixture
def docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = CELLS[0]
    table.cell(0, 1).text = CELLS[1]
    for text in BODY:
        document.add_paragraph(text)
    path = tmp_path / "main.docx"
    document.save(str(path))
    return path


def display_fields(
    original_text: str, replacement_text: str, line: str | None
) -> Dict[str, object]:
    """The rendered columns the reporter would have stored for this edit.

    Worked out the way `IssueReporter._resolve_edit` works them out: the quote
    is located on its markdown line and that line is rendered with the quote
    marked, so a test's edit carries exactly what a real one would. A quote
    the line does not carry -- a fixture checking a note's wording, or one
    deliberately absent from the document -- falls back to the quote itself,
    which is what an edit over plain prose renders to anyway.
    """
    spans = (
        locate_in_paragraph(line, normalize_whitespace(original_text))
        if line is not None
        else []
    )
    rendered = (
        rendered_span(line, spans[0][0], spans[0][1])
        if line is not None and len(spans) == 1
        else None
    )
    return {
        "display_text": (
            rendered.display_text if rendered else normalize_whitespace(original_text)
        ),
        "display_occurrence": rendered.display_occurrence if rendered else 0,
        "display_replacement": render_replacement_text(
            replacement_text,
            # As the reporter decides it: a quote opening its line took the
            # line's block syntax with it, and so did its replacement.
            block_context=bool(rendered) and spans[0][0] == 0,
        ),
    }


def make_edit(
    original_text: str,
    replacement_text: str,
    start_line: int,
    *,
    status: IssueEditStatus = IssueEditStatus.PROPOSED,
    edit_id: uuid.UUID | None = None,
) -> IssueEdit:
    """One proposed-edit row, with its rendered columns filled from `MARKDOWN`.

    `plan_edits` fills them again from whatever markdown the test actually
    plans against, so a suite with its own document does not have to repeat it
    here.
    """
    lines = document_lines(MARKDOWN)
    line = lines[start_line - 1] if 1 <= start_line <= len(lines) else None
    return IssueEdit(
        id=edit_id or uuid.uuid4(),
        issue_id=uuid.uuid4(),
        position=0,
        original_text=original_text,
        replacement_text=replacement_text,
        start_line=start_line,
        end_line=start_line,
        rationale="because the figure is misnumbered",
        status=status,
        **display_fields(original_text, replacement_text, line),
    )


def make_candidates(
    edits: Sequence[IssueEdit], severities: Sequence[SeverityEnum] | None = None
) -> List[EditCandidate]:
    chosen = severities or [SeverityEnum.MEDIUM] * len(edits)
    return [
        EditCandidate(
            edit=edit,
            issue_severity=severity,
            issue_created_at=datetime(2026, 1, 1, tzinfo=UTC),
            workflow_type="reference_validation_v2",
        )
        for edit, severity in zip(edits, chosen)
    ]


async def plan_edits(
    path: Path,
    edits: Sequence[IssueEdit],
    severities: Sequence[SeverityEnum] | None = None,
    paragraph_line_ranges: Dict[int, Tuple[int, int]] | None = None,
    markdown: str | None = None,
):
    candidates = make_candidates(edits, severities)
    lines = document_lines(MARKDOWN if markdown is None else markdown)
    # The reporter derives the rendered columns from the document it read; a
    # test planning against its own markdown gets the same treatment here,
    # rather than every call site repeating the document twice.
    for edit in edits:
        source = (
            lines[edit.start_line - 1] if 1 <= edit.start_line <= len(lines) else None
        )
        for name, value in display_fields(
            edit.original_text, edit.replacement_text, source
        ).items():
            setattr(edit, name, value)
    decisions = resolve_edit_conflicts(candidates, lines)
    plan = await plan_tracked_changes(
        str(path),
        decisions,
        {candidate.edit.id: candidate for candidate in candidates},
        (
            PARAGRAPH_LINE_RANGES
            if paragraph_line_ranges is None
            else paragraph_line_ranges
        ),
        lines,
        workspace_root=str(path.parent),
    )
    return plan, decisions


def statuses(outcomes: Sequence[EditOutcome]) -> Dict[uuid.UUID, str]:
    return {outcome.edit_id: outcome.status for outcome in outcomes}


def read_document(path: Path):
    doc = EditorDocument.open(
        path, author="Reader", workspace_dir=str(path.parent / "read-ws")
    )
    try:
        return doc.get_visible_text(), doc.get_original_text(), doc.list_revisions()
    finally:
        doc.close()
