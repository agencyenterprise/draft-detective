"""Tests for writing the planned redlines into the file.

What the output still reads as afterwards, what an existing redline does to the
paragraph mapping, and what happens when the two libraries disagree about the
document at all.
"""

from pathlib import Path
from typing import Dict, Sequence, Tuple

import pytest
from docx import Document as PythonDocxDocument
from docx_editor import Document as EditorDocument, EditOperation

from lib.services.docx.tracked_changes import apply_tracked_changes

from tests.unit.services.tracked_changes_support import (
    BODY as _BODY,
    OTHER_PASSAGE_DETAIL as _OTHER_PASSAGE_DETAIL,
    docx_path,  # noqa: F401 -- a fixture, reached by name
    make_edit as _edit,
    plan_edits as _plan,
    read_document as _read,
    statuses as _statuses,
)


class TestTheFileStaysReadable:
    @pytest.mark.asyncio
    async def test_python_docx_reopens_the_output_with_untouched_text_intact(
        self, docx_path: Path
    ):
        edits = [
            _edit("14% rise", "18% rise", 1),
            _edit("**Figure 3**", "Figure 4", 3),
        ]

        plan, _ = await _plan(docx_path, edits)
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        paragraphs = [
            p.text
            for p in PythonDocxDocument(str(docx_path)).paragraphs
            if p.text.strip()
        ]
        # The redlined paragraphs no longer read as plain text to python-docx --
        # which is exactly why the comment pass has to run first -- but the
        # untouched one is unchanged.
        assert _BODY[2] in paragraphs

    @pytest.mark.asyncio
    async def test_an_empty_plan_leaves_the_file_alone(self, docx_path: Path):
        before = docx_path.read_bytes()

        assert await apply_tracked_changes(str(docx_path), []) == []
        assert docx_path.read_bytes() == before


# Two paragraphs reading alike, the first of which already carries someone
# else's tracked insertion. python-docx leaves an inserted run out of
# `paragraph.text`, so the two are indistinguishable by text on its side while
# docx-editor shows the insertion -- the case that made the old text-based
# mapping hand the redline to the wrong paragraph.
_ALIKE = "The results are significant."
_ALIKE_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 1), 1: (2, 2)}


@pytest.fixture
def alike_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph(_ALIKE)
    document.add_paragraph(_ALIKE)
    path = tmp_path / "alike.docx"
    document.save(str(path))

    doc = EditorDocument.open(
        path, author="Someone else", workspace_dir=str(tmp_path / "prior-ws")
    )
    try:
        first = doc.list_paragraphs_structured(limit=None)[0]
        doc.batch_edit(
            [
                EditOperation.replace(
                    "The results", "New The results", paragraph=first.ref
                )
            ]
        )
        doc.save()
    finally:
        doc.close()
    return path


class TestAnExistingRedlineDoesNotMisdirectTheNewOne:
    @pytest.mark.asyncio
    async def test_the_edit_lands_on_the_paragraph_its_line_names(
        self, alike_docx_path: Path
    ):
        edit = _edit("results", "findings", 1)

        plan, _ = await _plan(
            alike_docx_path,
            [edit],
            paragraph_line_ranges=_ALIKE_LINE_RANGES,
            markdown="\n".join([_ALIKE, _ALIKE]),
        )
        await apply_tracked_changes(
            str(alike_docx_path),
            plan.planned,
            workspace_root=str(alike_docx_path.parent),
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        # Paragraph 1 is the one that carried the earlier insertion.
        doc = EditorDocument.open(
            alike_docx_path,
            author="Reader",
            workspace_dir=str(alike_docx_path.parent / "check-ws"),
        )
        try:
            texts = [info.text for info in doc.list_paragraphs_structured(limit=None)]
        finally:
            doc.close()
        assert texts == ["New The findings are significant.", _ALIKE]


class TestParagraphsThatCannotBeMatched:
    @pytest.mark.asyncio
    async def test_disagreeing_paragraph_counts_place_nothing(
        self, docx_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        original = EditorDocument.list_paragraphs_structured

        def one_short(self, **kwargs):  # type: ignore[no-untyped-def]
            return original(self, **kwargs)[:-1]

        monkeypatch.setattr(EditorDocument, "list_paragraphs_structured", one_short)
        edit = _edit("14% rise", "18% rise", 1)

        plan, _ = await _plan(docx_path, [edit])
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "not_found"}
        assert [outcome.detail for outcome in plan.outcomes] == [
            "the document's paragraphs could not be matched to the file being "
            "exported"
        ]


# One Word paragraph holding hard line breaks, which MarkItDown converts to one
# markdown line per break. Word keeps nothing where a break was, so the
# paragraph reads as its lines run together and a quote's line-relative
# occurrence says nothing about the paragraph's own spans.
_BREAK_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 2)}
_THREE_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 3)}

# Why a repeat in such a paragraph is reported rather than placed.
_REPEATS_DETAIL = (
    "the quoted text repeats inside a paragraph that spans several lines "
    "of the document"
)


def _hard_break_docx(path: Path, runs: Sequence[str]) -> Path:
    """One paragraph whose runs are separated by hard line breaks."""
    document = PythonDocxDocument()
    paragraph = document.add_paragraph()
    for position, text in enumerate(runs):
        run = paragraph.add_run(text)
        if position < len(runs) - 1:
            run.add_break()
    document.save(str(path))
    return path


class TestAParagraphSpanningSeveralMarkdownLines:
    """A quote the paragraph carries once is placed; a repeat is reported.

    `display_occurrence` was counted on the edit's own markdown line, and this
    paragraph is several of them, so the index does not apply. Rebasing it
    proved wrong in three different ways -- a match straddling the join, a line
    an earlier line contains, boundaries lost to conversion drift -- so a
    repeat is left to a human.
    """

    async def _plan_break(
        self,
        path: Path,
        runs: Sequence[str],
        markdown_lines: Sequence[str],
        quote: str,
        replacement: str,
        line: int,
        ranges: Dict[int, Tuple[int, int]] | None = None,
    ):
        # The bytes of this save, for tests asserting the file is untouched.
        # Comparing against an earlier save is flaky: a zip entry records its
        # save time at two-second resolution.
        saved = _hard_break_docx(path, runs).read_bytes()
        edit = _edit(quote, replacement, line)
        plan, _ = await _plan(
            path,
            [edit],
            paragraph_line_ranges=ranges or _BREAK_LINE_RANGES,
            markdown="\n".join(markdown_lines),
        )
        await apply_tracked_changes(
            str(path), plan.planned, workspace_root=str(path.parent)
        )
        return edit, plan, saved

    @pytest.mark.asyncio
    async def test_a_quote_the_paragraph_carries_once_is_placed(self, tmp_path: Path):
        runs = ["First cohort rose 14%.", "Second cohort rose 9%."]
        edit, plan, _ = await self._plan_break(
            tmp_path / "unique.docx", runs, runs, "9%", "18%", 2
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        visible, original, _ = _read(tmp_path / "unique.docx")
        assert "First cohort rose 14%." in visible
        assert "Second cohort rose 18%." in visible
        assert "Second cohort rose 9%." in original

    @pytest.mark.asyncio
    async def test_a_plain_repeat_across_the_break_is_reported(self, tmp_path: Path):
        runs = ["First cohort rose 14%.", "Second cohort rose 14%."]
        path = tmp_path / "repeat.docx"
        edit, plan, before = await self._plan_break(path, runs, runs, "14%", "18%", 2)

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "ambiguous"}
        assert [outcome.detail for outcome in plan.outcomes] == [_REPEATS_DETAIL]
        assert path.read_bytes() == before
        _, _, revisions = _read(path)
        assert list(revisions) == []

    @pytest.mark.asyncio
    async def test_a_repeat_the_join_itself_created_is_reported(self, tmp_path: Path):
        # A space before the break puts a "14% in" in the paragraph that
        # neither markdown line carries.
        runs = ["First cohort rose 14% ", "in 2019; second cohort rose 14% in 2020."]
        markdown = ["First cohort rose 14%", runs[1]]
        path = tmp_path / "straddle.docx"
        edit, plan, before = await self._plan_break(
            path, runs, markdown, "14% in", "18% in", 2
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "ambiguous"}
        assert [outcome.detail for outcome in plan.outcomes] == [_REPEATS_DETAIL]
        assert path.read_bytes() == before

    @pytest.mark.asyncio
    async def test_a_repeat_across_two_identical_lines_is_reported(
        self, tmp_path: Path
    ):
        runs = ["The first cohort rose 14%.", "The first cohort rose 14%."]
        path = tmp_path / "twin.docx"
        edit, plan, before = await self._plan_break(path, runs, runs, "14%", "18%", 2)

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "ambiguous"}
        assert [outcome.detail for outcome in plan.outcomes] == [_REPEATS_DETAIL]
        assert path.read_bytes() == before

    @pytest.mark.asyncio
    async def test_a_repeat_inside_a_line_an_earlier_line_contains_is_reported(
        self, tmp_path: Path
    ):
        # Line 2 reads as the tail of line 1, so the paragraph shows the
        # figure twice.
        runs = [
            "Group A reported a sustained 14% increase.",
            "A sustained 14% increase.",
        ]
        path = tmp_path / "nested.docx"
        edit, plan, before = await self._plan_break(path, runs, runs, "14%", "18%", 2)

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "ambiguous"}
        assert [outcome.detail for outcome in plan.outcomes] == [_REPEATS_DETAIL]
        assert path.read_bytes() == before

    @pytest.mark.asyncio
    async def test_a_repeat_three_lines_apart_is_reported(self, tmp_path: Path):
        runs = [
            "The first cohort rose 14%.",
            "The second cohort fell 9%.",
            "The third cohort rose 14%.",
        ]
        path = tmp_path / "reprise.docx"
        edit, plan, before = await self._plan_break(
            path, runs, runs, "14%", "18%", 3, ranges=_THREE_LINE_RANGES
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "ambiguous"}
        assert [outcome.detail for outcome in plan.outcomes] == [_REPEATS_DETAIL]
        assert path.read_bytes() == before


# An ordinary one-paragraph-per-line document, with the blank separator the
# converter writes. The mapper hands the paragraph the range up to the line
# before the next one, so its range covers the separator -- which must not make
# the paragraph "multi-line", or `display_occurrence` would be thrown away on
# every paragraph in the document.
_SEPARATED_MARKDOWN = "\n".join(["Figure 3 and **Figure 3** close the section.", ""])
_SEPARATED_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 2)}


@pytest.fixture
def separated_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph("Figure 3 and Figure 3 close the section.")
    path = tmp_path / "separated.docx"
    document.save(str(path))
    return path


class TestABlankSeparatorIsNotAnotherLine:
    @pytest.mark.asyncio
    async def test_a_repeat_still_resolves_by_the_stored_occurrence(
        self, separated_docx_path: Path
    ):
        edit = _edit("**Figure 3**", "Figure 4", 1)

        plan, _ = await _plan(
            separated_docx_path,
            [edit],
            paragraph_line_ranges=_SEPARATED_LINE_RANGES,
            markdown=_SEPARATED_MARKDOWN,
        )
        await apply_tracked_changes(
            str(separated_docx_path),
            plan.planned,
            workspace_root=str(separated_docx_path.parent),
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        # The bold repeat, not the plain one before it.
        assert plan.planned[0].occurrence == 1
        visible, _, _ = _read(separated_docx_path)
        assert "Figure 3 and Figure 4 close the section." in visible


# A nested line short enough to sit inside the paragraph above it by accident:
# a list item or a cell rendering to just a percentage is contained by any
# paragraph that quotes one.
_SHORT_NESTED_PARAGRAPH = "Output for the western region increased by 14%."
_SHORT_NESTED_MARKDOWN = "\n".join([_SHORT_NESTED_PARAGRAPH, "", "- 14%"])
_SHORT_NESTED_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 3)}


@pytest.fixture
def short_nested_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph(_SHORT_NESTED_PARAGRAPH)
    path = tmp_path / "short-nested.docx"
    document.save(str(path))
    return path


class TestAShortNestedLineIsNotTheParagraph:
    @pytest.mark.asyncio
    async def test_a_line_the_paragraph_merely_contains_is_rejected(
        self, short_nested_docx_path: Path
    ):
        edit = _edit("14%", "18%", 3)

        plan, _ = await _plan(
            short_nested_docx_path,
            [edit],
            paragraph_line_ranges=_SHORT_NESTED_RANGES,
            markdown=_SHORT_NESTED_MARKDOWN,
        )
        await apply_tracked_changes(
            str(short_nested_docx_path),
            plan.planned,
            workspace_root=str(short_nested_docx_path.parent),
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "unlocatable"}
        assert [outcome.detail for outcome in plan.outcomes] == [_OTHER_PASSAGE_DETAIL]
        visible, _, revisions = _read(short_nested_docx_path)
        assert visible.strip() == _SHORT_NESTED_PARAGRAPH
        assert list(revisions) == []

    @pytest.mark.asyncio
    async def test_a_line_long_enough_to_be_contained_on_purpose_still_passes(
        self, tmp_path: Path
    ):
        # 25 characters, contained in a hard-break paragraph that really does
        # carry it.
        runs = ["Group A reported a sustained rise.", "A sustained 14% increase."]
        path = _hard_break_docx(tmp_path / "long-contained.docx", runs)
        edit = _edit("14%", "18%", 2)

        plan, _ = await _plan(
            path,
            [edit],
            paragraph_line_ranges=_BREAK_LINE_RANGES,
            markdown="\n".join(runs),
        )
        await apply_tracked_changes(
            str(path), plan.planned, workspace_root=str(path.parent)
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        visible, _, _ = _read(path)
        assert "A sustained 18% increase." in visible

    @pytest.mark.asyncio
    async def test_a_short_line_that_is_the_whole_paragraph_still_passes(
        self, tmp_path: Path
    ):
        # Equality needs no length floor: the line *is* the paragraph.
        document = PythonDocxDocument()
        document.add_paragraph("Yield 14%.")
        path = tmp_path / "short-equal.docx"
        document.save(str(path))
        edit = _edit("14%", "18%", 1)

        plan, _ = await _plan(
            path,
            [edit],
            paragraph_line_ranges={0: (1, 1)},
            markdown="Yield 14%.",
        )
        await apply_tracked_changes(
            str(path), plan.planned, workspace_root=str(path.parent)
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        visible, _, _ = _read(path)
        assert "Yield 18%." in visible


# An underscore-delimited fragment of a URL: the mark the renderer inserts at
# the quote's boundary must not turn the intraword `_` into emphasis, or the
# export searches the paragraph for `final` and leaves the underscores behind.
_URL_PARAGRAPH = "See https://example.org/report_final_version.pdf for details."


@pytest.fixture
def url_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph(_URL_PARAGRAPH)
    path = tmp_path / "url.docx"
    document.save(str(path))
    return path


class TestAnIntrawordUnderscoreIsNotEmphasis:
    @pytest.mark.asyncio
    async def test_the_redline_covers_the_underscores_the_quote_named(
        self, url_docx_path: Path
    ):
        edit = _edit("_final_", "-draft-", 1)

        plan, _ = await _plan(
            url_docx_path,
            [edit],
            paragraph_line_ranges={0: (1, 1)},
            markdown=_URL_PARAGRAPH,
        )
        await apply_tracked_changes(
            str(url_docx_path), plan.planned, workspace_root=str(url_docx_path.parent)
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        assert plan.planned[0].find == "_final_"
        assert plan.planned[0].replace_with == "-draft-"
        visible, original, _ = _read(url_docx_path)
        assert "report-draft-version.pdf" in visible
        assert "report_final_version.pdf" in original
