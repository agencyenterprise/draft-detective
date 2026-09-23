"""Blank python-docx documents without the default template's style bulk.

python-docx starts every new document from a template carrying ~790KB of style
XML (`styles.xml` plus `stylesWithEffects.xml`). docx-editor pretty-prints and
re-parses every part on each open and save, so a test opening a two-paragraph
document a few times spent seconds on styles it never uses. The lean template
keeps only the default styles, which is all an unstyled paragraph or table
resolves to; the document body is built exactly as before.
"""

import io
import zipfile
from functools import cache
from typing import Callable

from docx import Document
from docx.document import Document as DocxDocument
from lxml import etree

_STYLES_WITH_EFFECTS = "stylesWithEffects.xml"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_REFERENCING_PARTS = ("[Content_Types].xml", "word/_rels/document.xml.rels")


def new_document() -> DocxDocument:
    """A blank document, as `docx.Document()` would give, minus unused styles."""
    return Document(io.BytesIO(_lean_template()))


@cache
def _lean_template() -> bytes:
    source = io.BytesIO()
    Document().save(source)
    lean = io.BytesIO()
    with (
        zipfile.ZipFile(source) as original,
        zipfile.ZipFile(lean, "w", zipfile.ZIP_DEFLATED) as trimmed,
    ):
        for info in original.infolist():
            if info.filename.endswith(_STYLES_WITH_EFFECTS):
                continue
            trimmed.writestr(info, _trim_part(info.filename, original.read(info)))
    return lean.getvalue()


def _trim_part(name: str, data: bytes) -> bytes:
    if name == "word/styles.xml":
        return _without_children(data, _is_unused_style)
    if name in _REFERENCING_PARTS:
        return _without_children(data, _references_styles_with_effects)
    return data


def _is_unused_style(element: etree._Element) -> bool:
    if element.tag == f"{_W}latentStyles":
        return True
    return element.tag == f"{_W}style" and element.get(f"{_W}default") != "1"


def _references_styles_with_effects(element: etree._Element) -> bool:
    target = element.get("PartName") or element.get("Target") or ""
    return target.endswith(_STYLES_WITH_EFFECTS)


def _without_children(data: bytes, drop: Callable[[etree._Element], bool]) -> bytes:
    root = etree.fromstring(data)
    for child in list(root):
        if drop(child):
            root.remove(child)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
