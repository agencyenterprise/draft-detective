"""Tests for tying a python-docx paragraph index to a docx-editor ref.

Built on a real .docx and read back with both libraries, because the whole
point of the mapping is that the two enumerate a document the same way.
"""

import copy
from pathlib import Path
from typing import Any, List

import pytest
from docx import Document as PythonDocxDocument
from docx.oxml.ns import qn
from docx_editor import Document as EditorDocument, EditOperation, ParagraphInfo
from lxml import etree

from lib.services.docx.paragraph_refs import (
    map_paragraph_refs,
    paragraph_ordinals,
)

from tests.unit.services.lean_docx import new_document

# A body paragraph, a two-cell table, an empty paragraph, two paragraphs that
# read alike, and a paragraph inside a content control: everything that makes
# the two libraries' paragraph lists differ.
_SHARED = "The results are significant."


@pytest.fixture
def mixed_docx_path(tmp_path: Path) -> Path:
    document = new_document()
    document.add_paragraph("First body paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell one"
    table.cell(0, 1).text = "Cell two"
    document.add_paragraph("")
    document.add_paragraph(_SHARED)
    document.add_paragraph(_SHARED)
    _append_content_control(document, "Inside a content control.")
    path = tmp_path / "mixed.docx"
    document.save(str(path))
    return path


def _append_content_control(document: Any, text: str) -> None:
    """Wrap a copy of the last paragraph in a ``w:sdt``, carrying `text`.

    `document` is a python-docx ``Document``, which the library exposes as a
    factory function rather than a type.

    python-docx cannot build one, and a content control is the other structure
    that puts a ``<w:p>`` outside ``Document.paragraphs``.
    """
    sdt = etree.SubElement(document.element.body, qn("w:sdt"))
    etree.SubElement(sdt, qn("w:sdtPr"))
    content = etree.SubElement(sdt, qn("w:sdtContent"))
    paragraph = copy.deepcopy(document.paragraphs[-1]._p)
    for node in paragraph.iter(qn("w:t")):
        node.text = text
    content.append(paragraph)


def _editor_paragraphs(path: Path) -> List[ParagraphInfo]:
    doc = EditorDocument.open(
        path, author="Reader", workspace_dir=str(path.parent / "read-ws")
    )
    try:
        return doc.list_paragraphs_structured(limit=None)
    finally:
        doc.close()


class TestParagraphOrdinals:
    def test_every_paragraph_of_the_file_is_counted(self, mixed_docx_path: Path):
        walk = paragraph_ordinals(PythonDocxDocument(str(mixed_docx_path)))

        assert walk is not None
        ordinals, total = walk
        # Body paragraph, two cells, the empty one, two alike, the control.
        assert total == 7
        assert total == len(_editor_paragraphs(mixed_docx_path))

    def test_only_the_non_empty_body_paragraphs_are_mapped(self, mixed_docx_path: Path):
        walk = paragraph_ordinals(PythonDocxDocument(str(mixed_docx_path)))

        assert walk is not None
        ordinals, _ = walk
        # The table cells (2, 3), the empty paragraph (4) and the content
        # control (7) are not paragraphs the comment pass indexes.
        assert ordinals == [1, 5, 6]


class TestMapParagraphRefs:
    def test_each_index_gets_the_ref_at_its_own_position(self, mixed_docx_path: Path):
        editor_paragraphs = _editor_paragraphs(mixed_docx_path)
        walk = paragraph_ordinals(PythonDocxDocument(str(mixed_docx_path)))
        assert walk is not None

        mapped = map_paragraph_refs(walk[0], walk[1], editor_paragraphs)

        assert mapped is not None
        assert [mapped[i].ref for i in sorted(mapped)] == [
            editor_paragraphs[0].ref,
            editor_paragraphs[4].ref,
            editor_paragraphs[5].ref,
        ]
        assert [mapped[i].text for i in sorted(mapped)] == [
            "First body paragraph.",
            _SHARED,
            _SHARED,
        ]

    def test_an_existing_insertion_does_not_move_the_mapping(
        self, mixed_docx_path: Path, tmp_path: Path
    ):
        # An insertion someone else left in the first of the two alike
        # paragraphs. python-docx does not read it as paragraph text, so
        # matching by text would hand index 1 to the other paragraph.
        _insert_tracked_text(mixed_docx_path, tmp_path)
        editor_paragraphs = _editor_paragraphs(mixed_docx_path)
        walk = paragraph_ordinals(PythonDocxDocument(str(mixed_docx_path)))
        assert walk is not None

        mapped = map_paragraph_refs(walk[0], walk[1], editor_paragraphs)

        assert mapped is not None
        assert mapped[1].text == "The new results are significant."
        assert mapped[2].text == _SHARED

    def test_disagreeing_counts_map_nothing(self, mixed_docx_path: Path):
        editor_paragraphs = _editor_paragraphs(mixed_docx_path)

        assert map_paragraph_refs([1, 5, 6], 7, editor_paragraphs[:-1]) is None

    def test_a_missing_paragraph_maps_nothing(self, mixed_docx_path: Path):
        editor_paragraphs = _editor_paragraphs(mixed_docx_path)

        # The counts agree, but no paragraph carries ordinal 9.
        assert (
            map_paragraph_refs([9], len(editor_paragraphs), editor_paragraphs) is None
        )

    def test_a_document_without_paragraphs_maps_nothing_and_raises_nothing(self):
        assert map_paragraph_refs([], 0, []) == {}


def _insert_tracked_text(path: Path, tmp_path: Path) -> None:
    """Leave a tracked insertion in the first paragraph reading `_SHARED`."""
    doc = EditorDocument.open(
        path, author="Someone else", workspace_dir=str(tmp_path / "insert-ws")
    )
    try:
        target = next(
            info
            for info in doc.list_paragraphs_structured(limit=None)
            if info.text == _SHARED
        )
        doc.batch_edit(
            [
                EditOperation.replace(
                    "The results", "The new results", paragraph=target.ref
                )
            ]
        )
        doc.save()
    finally:
        doc.close()
