"""Tests for treating a Flat OPC package as an editable document.

The adapter's job is to lose nothing. A part dropped in transit comes back from
Word as ``GeneralException`` with no detail, so the tests here are mostly about
what survives the round trip rather than what the library does with it.
"""

import base64
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document as NewDocument
from lxml import etree

from lib.services.docx.namespaces import PKG, qn
from lib.services.microsoft.word.word_package import (
    DRAFT_DETECTIVE,
    FragmentError,
    editable,
    from_docx,
    to_docx,
)


def flat_opc(tmp_path: Path, *paragraphs: str) -> str:
    """A package shaped like the one an add-in sends, built from a real .docx."""

    document = NewDocument()
    for text in paragraphs or ("A paragraph.",):
        document.add_paragraph(text)
    source = tmp_path / "built.docx"
    document.save(str(source))
    return from_docx(source)


def part_names(markup: str) -> list[str]:
    root = etree.fromstring(markup.encode("utf-8"))
    return [
        node.get(qn("name", PKG)) or "" for node in root.iterfind(qn("part", PKG))
    ]


class TestTheRoundTrip:
    def test_every_part_survives(self, tmp_path: Path) -> None:
        """Everything but the content-types part, which is rebuilt rather than carried."""

        markup = flat_opc(tmp_path)
        rebuilt = tmp_path / "rebuilt.docx"
        to_docx(markup, rebuilt)

        with ZipFile(rebuilt) as archive:
            entries = {name for name in archive.namelist() if not name.endswith("/")}

        assert "[Content_Types].xml" in entries, "a .docx is invalid without it"
        assert {"/" + name for name in entries if name != "[Content_Types].xml"} == set(
            part_names(markup)
        )

    def test_the_rebuilt_content_types_names_every_part(self, tmp_path: Path) -> None:
        """A part the content types do not mention is a part no reader will accept."""

        markup = flat_opc(tmp_path)
        rebuilt = tmp_path / "rebuilt.docx"
        to_docx(markup, rebuilt)

        with ZipFile(rebuilt) as archive:
            declaration = archive.read("[Content_Types].xml").decode()
            entries = [
                name
                for name in archive.namelist()
                if name.endswith(".xml") and name != "[Content_Types].xml"
            ]
        for name in entries:
            assert f'PartName="/{name}"' in declaration, f"{name} is undeclared"

    def test_word_keeps_its_prolog(self, tmp_path: Path) -> None:
        """Without it the package is not recognised as a Word document."""

        assert "mso-application" in flat_opc(tmp_path)

    def test_content_types_come_from_the_package(self, tmp_path: Path) -> None:
        markup = flat_opc(tmp_path)
        root = etree.fromstring(markup.encode("utf-8"))
        types = {
            node.get(qn("name", PKG)): node.get(qn("contentType", PKG))
            for node in root.iterfind(qn("part", PKG))
        }

        assert types["/word/document.xml"] == (
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document.main+xml"
        )
        assert types["/_rels/.rels"] == (
            "application/vnd.openxmlformats-package.relationships+xml"
        )

    def test_a_binary_part_is_not_mangled(self, tmp_path: Path) -> None:
        """Images travel as base64; decoding them badly corrupts the document."""

        markup = flat_opc(tmp_path)
        payload = b"\x89PNG\r\n\x1a\n" + bytes(range(256))
        root = etree.fromstring(markup.encode("utf-8"))
        part = etree.SubElement(root, qn("part", PKG))
        part.set(qn("name", PKG), "/word/media/image1.png")
        part.set(qn("contentType", PKG), "image/png")
        binary = etree.SubElement(part, qn("binaryData", PKG))
        binary.text = base64.b64encode(payload).decode("ascii")

        rebuilt = tmp_path / "with-image.docx"
        to_docx(etree.tostring(root, encoding="unicode"), rebuilt)
        with ZipFile(rebuilt) as archive:
            assert archive.read("word/media/image1.png") == payload

        # and back out again, still intact
        again = from_docx(rebuilt)
        assert "/word/media/image1.png" in part_names(again)
        restored = etree.fromstring(again.encode("utf-8"))
        for node in restored.iterfind(qn("part", PKG)):
            if node.get(qn("name", PKG)) == "/word/media/image1.png":
                data = node.find(qn("binaryData", PKG))
                assert data is not None and data.text
                assert base64.b64decode(data.text) == payload
                break
        else:
            pytest.fail("the image part did not come back")

    def test_no_part_carries_a_second_xml_declaration(self, tmp_path: Path) -> None:
        """A declaration inside the package would make it malformed."""

        markup = flat_opc(tmp_path)
        assert markup.count("<?xml") == 1


class TestWhatWordActuallySends:
    """Regressions from a real Word document, every one of which reached Word.

    The earlier tests here all built fixtures with python-docx, whose output is
    tidier than Word's in exactly the ways that mattered. Each test below stands
    for a package Word rejected with ``GeneralException`` and no detail.
    """

    def word_like_docx(self, tmp_path: Path) -> Path:
        """A .docx with the awkward parts a synced Word file really carries."""

        base = tmp_path / "base.docx"
        document = NewDocument()
        document.add_paragraph("Testing")
        document.save(str(base))

        out = tmp_path / "wordlike.docx"
        with ZipFile(base) as source, ZipFile(out, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    declaration = payload.split(b"?>", 1)
                    body = declaration[1] if len(declaration) > 1 else payload
                    payload = (
                        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + body
                    )
                    # Word writes a byte-order mark; str.lstrip() ignores it.
                    payload = b"\xef\xbb\xbf" + payload
                target.writestr(info.filename, payload)
            # OneDrive deletion leftovers: no relationship, no declared type.
            target.writestr("[trash]/0000.dat", b"\x00\x01leftover")
            # SharePoint custom XML opens with a processing instruction.
            target.writestr(
                "customXml/item2.xml",
                b'<?mso-contentType?><FormTemplates xmlns="urn:x">'
                b"<Display>Form</Display></FormTemplates>",
            )
        return out

    def test_a_byte_order_mark_does_not_nest_a_declaration(
        self, tmp_path: Path
    ) -> None:
        """A second declaration inside the package makes the whole thing unparseable."""

        markup = from_docx(self.word_like_docx(tmp_path))

        assert markup.count("<?xml") == 1
        etree.fromstring(markup.encode("utf-8"))  # would raise if malformed

    def test_a_part_with_no_declared_type_is_left_out(self, tmp_path: Path) -> None:
        """Inventing a type for [trash] entries is what Word rejected."""

        markup = from_docx(self.word_like_docx(tmp_path))
        assert not any("trash" in name for name in part_names(markup))

    def test_the_content_types_part_is_not_carried(self, tmp_path: Path) -> None:
        """Word's own getOoxml ships none; the types live on each part instead."""

        markup = from_docx(self.word_like_docx(tmp_path))
        assert "/[Content_Types].xml" not in part_names(markup)

    def test_a_package_without_content_types_still_opens(self, tmp_path: Path) -> None:
        """The shape Word actually sends. to_docx has to rebuild the part."""

        markup = from_docx(self.word_like_docx(tmp_path))
        assert "/[Content_Types].xml" not in part_names(markup)

        rebuilt = tmp_path / "rebuilt.docx"
        to_docx(markup, rebuilt)
        with ZipFile(rebuilt) as archive:
            assert "[Content_Types].xml" in archive.namelist()

        with editable(markup) as fragment:
            assert fragment.document.list_paragraphs_structured()[0].text == "Testing"

    def test_a_part_beginning_with_an_instruction_keeps_its_root(
        self, tmp_path: Path
    ) -> None:
        """Taking the first child wrote out the instruction and dropped the document."""

        markup = from_docx(self.word_like_docx(tmp_path))
        rebuilt = tmp_path / "rebuilt.docx"
        to_docx(markup, rebuilt)

        with ZipFile(rebuilt) as archive:
            payload = archive.read("customXml/item2.xml")
        assert b"FormTemplates" in payload, "the root element must survive"
        assert b"mso-contentType" in payload, "and so must what preceded it"
        etree.fromstring(payload)

    def test_the_whole_round_trip_survives_a_word_like_package(
        self, tmp_path: Path
    ) -> None:
        """Read, edit, write back: what the routes do on every request."""

        markup = from_docx(self.word_like_docx(tmp_path))
        with editable(markup) as fragment:
            info = fragment.document.list_paragraphs_structured()[0]
            fragment.document.replace("Testing", "TESTED", paragraph=info.ref)
            out = fragment.serialise()

        etree.fromstring(out.encode("utf-8"))
        assert f'w:author="{DRAFT_DETECTIVE}"' in out
        assert not any("trash" in name for name in part_names(out))
        assert "/[Content_Types].xml" not in part_names(out)


class TestRefusals:
    def test_a_package_without_a_document_is_refused(self, tmp_path: Path) -> None:
        markup = (
            f'<pkg:package xmlns:pkg="{PKG}">'
            '<pkg:part pkg:name="/word/styles.xml" pkg:contentType="application/xml">'
            "<pkg:xmlData><w:styles/></pkg:xmlData></pkg:part></pkg:package>"
        ).replace("<w:styles/>", '<styles xmlns="urn:x"/>')
        with pytest.raises(FragmentError, match="not a document"):
            to_docx(markup, tmp_path / "broken.docx")

    def test_markup_that_is_not_a_package_is_refused(self, tmp_path: Path) -> None:
        with pytest.raises(Exception):
            to_docx("not xml at all", tmp_path / "broken.docx")


class TestEditing:
    def test_a_tracked_change_is_attributed_to_draft_detective(
        self, tmp_path: Path
    ) -> None:
        markup = flat_opc(tmp_path, "The results clearly prove the point.")

        with editable(markup) as fragment:
            ref = fragment.document.list_paragraphs_structured()[0].ref
            fragment.document.replace("clearly prove", "suggest", paragraph=ref)
            out = fragment.serialise()

        assert f'w:author="{DRAFT_DETECTIVE}"' in out
        assert "<w:ins " in out and "<w:del " in out
        assert "<w:delText" in out

    def test_a_comment_is_attributed_and_its_parts_are_written(
        self, tmp_path: Path
    ) -> None:
        """The registration parts are what Word rejects a package for missing."""

        markup = flat_opc(tmp_path, "Smith et al. report a 42% reduction.")

        with editable(markup) as fragment:
            fragment.document.add_comment("42% reduction", "Against which baseline?")
            out = fragment.serialise()

        names = part_names(out)
        for part in (
            "/word/comments.xml",
            "/word/commentsExtended.xml",
            "/word/commentsIds.xml",
            "/word/commentsExtensible.xml",
        ):
            assert part in names, f"{part} missing; Word rejects the package without it"
        assert DRAFT_DETECTIVE in out

    def test_the_workspace_is_cleaned_up(self, tmp_path: Path) -> None:
        markup = flat_opc(tmp_path)
        with editable(markup) as fragment:
            work = fragment._work
            assert work.exists()
        assert not work.exists(), "a workspace left behind leaks a lock and a temp dir"

    def test_serialising_twice_does_not_fail(self, tmp_path: Path) -> None:
        """Saving over its own source is something the library checks for."""

        markup = flat_opc(tmp_path, "One sentence here.")
        with editable(markup) as fragment:
            first = fragment.serialise()
            second = fragment.serialise()
        assert part_names(first) == part_names(second)
