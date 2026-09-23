"""Tests for edits whose markdown line is a table row.

A row is Word cells, never one of the body paragraphs the mapper indexes, so an
edit anchored to one is reported rather than redlined onto the prose above it --
even when the cell repeats that prose word for word.
"""

from pathlib import Path
from typing import Dict, Tuple

import pytest
from docx import Document as PythonDocxDocument

from lib.services.docx.tracked_changes import apply_tracked_changes

from tests.unit.services.lean_docx import new_document
from tests.unit.services.tracked_changes_support import (
    make_edit as _edit,
    plan_edits as _plan,
    read_document as _read,
    statuses as _statuses,
)

# A paragraph followed by a table, with the paragraph's own figure repeated in
# a cell. The line-range mapper runs a body paragraph's range up to the line
# before the next body paragraph starts, so the table's markdown rows fall
# inside paragraph 1's range -- the ranges below are that map, by hand.
_TABLE_MARKDOWN = "\n".join(
    [
        "# Title",
        "",
        "Output increased by 14%.",
        "",
        "| Metric | Value |",
        "| --- | --- |",
        "| Growth | 14% |",
    ]
)
_TABLE_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 1), 1: (3, 7)}

# Why an edit is turned down: because its line is a table row at all, rather
# than because the line and the paragraph are different passages.
_IN_A_TABLE_DETAIL = "the edit sits in a table, which the export cannot redline"


@pytest.fixture
def table_docx_path(tmp_path: Path) -> Path:
    document = new_document()
    document.add_paragraph("Title")
    document.add_paragraph("Output increased by 14%.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Growth"
    table.cell(1, 1).text = "14%"
    path = tmp_path / "table.docx"
    document.save(str(path))
    return path


class TestATableLineNeverRedlinesTheParagraphAbove:
    @pytest.mark.asyncio
    async def test_a_table_cell_edit_is_reported_rather_than_misplaced(
        self, table_docx_path: Path
    ):
        edit = _edit("14%", "18%", 7)

        plan, _ = await _plan(
            table_docx_path,
            [edit],
            paragraph_line_ranges=_TABLE_LINE_RANGES,
            markdown=_TABLE_MARKDOWN,
        )
        await apply_tracked_changes(
            str(table_docx_path),
            plan.planned,
            workspace_root=str(table_docx_path.parent),
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "unlocatable"}
        assert [outcome.detail for outcome in plan.outcomes] == [_IN_A_TABLE_DETAIL]
        visible, _, revisions = _read(table_docx_path)
        assert "Output increased by 14%." in visible
        assert list(revisions) == []

    @pytest.mark.asyncio
    async def test_the_same_quote_on_the_paragraphs_own_line_is_applied(
        self, table_docx_path: Path
    ):
        edit = _edit("14%", "18%", 3)

        plan, _ = await _plan(
            table_docx_path,
            [edit],
            paragraph_line_ranges=_TABLE_LINE_RANGES,
            markdown=_TABLE_MARKDOWN,
        )
        await apply_tracked_changes(
            str(table_docx_path),
            plan.planned,
            workspace_root=str(table_docx_path.parent),
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        visible, _, revisions = _read(table_docx_path)
        assert "Output increased by 18%." in visible
        assert [r.type for r in revisions] == ["deletion", "insertion"]


# A cell repeating its section's prose word for word. Comparing the row's text
# with the paragraph's would call them the same passage, so the row has to be
# recognised as a row first.
_ECHO_PARAGRAPH = "Output increased by 14%."
_ECHO_MARKDOWN = "\n".join(
    [
        _ECHO_PARAGRAPH,
        "",
        f"| {_ECHO_PARAGRAPH} |",
    ]
)
_ECHO_LINE_RANGES: Dict[int, Tuple[int, int]] = {0: (1, 3)}


@pytest.fixture
def echo_table_docx_path(tmp_path: Path) -> Path:
    document = new_document()
    document.add_paragraph(_ECHO_PARAGRAPH)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = _ECHO_PARAGRAPH
    path = tmp_path / "echo-table.docx"
    document.save(str(path))
    return path


class TestACellRepeatingTheParagraphIsStillACell:
    @pytest.mark.asyncio
    async def test_the_row_is_reported_and_the_paragraph_left_alone(
        self, echo_table_docx_path: Path
    ):
        edit = _edit("14%", "18%", 3)

        plan, _ = await _plan(
            echo_table_docx_path,
            [edit],
            paragraph_line_ranges=_ECHO_LINE_RANGES,
            markdown=_ECHO_MARKDOWN,
        )
        await apply_tracked_changes(
            str(echo_table_docx_path),
            plan.planned,
            workspace_root=str(echo_table_docx_path.parent),
        )

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "unlocatable"}
        assert [outcome.detail for outcome in plan.outcomes] == [_IN_A_TABLE_DETAIL]
        visible, _, revisions = _read(echo_table_docx_path)
        assert visible.count(_ECHO_PARAGRAPH) == 2
        assert list(revisions) == []
        reopened = PythonDocxDocument(str(echo_table_docx_path))
        assert [p.text for p in reopened.paragraphs if p.text.strip()] == [
            _ECHO_PARAGRAPH
        ]
        assert reopened.tables[0].cell(0, 0).text == _ECHO_PARAGRAPH
