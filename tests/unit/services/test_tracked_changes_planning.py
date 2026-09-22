"""Tests for resolving a proposed edit to one redline.

Which paragraph an edit lands in, which occurrence of a repeated phrase it
means, and what the plan hands docx-editor to search for.
"""

from pathlib import Path

import pytest
from docx import Document as PythonDocxDocument

from lib.services.docx.tracked_changes import apply_tracked_changes

from tests.unit.services.tracked_changes_support import (
    docx_path,  # noqa: F401 -- a fixture, reached by name
    make_edit as _edit,
    plan_edits as _plan,
    read_document as _read,
    statuses as _statuses,
)


class TestPlanAndApply:
    @pytest.mark.asyncio
    async def test_a_located_edit_becomes_a_tracked_replacement(self, docx_path: Path):
        edit = _edit("14% rise", "18% rise", 1)

        plan, _ = await _plan(docx_path, [edit])
        outcomes = await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        assert _statuses(outcomes) == {edit.id: "applied"}
        visible, original, revisions = _read(docx_path)
        assert "reports a 18% rise in output" in visible
        assert "reports a 14% rise in output" in original
        # One deletion and one insertion, narrowed by the library to the part
        # of the quote that actually changes.
        assert [r.type for r in revisions] == ["deletion", "insertion"]
        assert revisions[0].text in "14% rise"
        assert revisions[1].text in "18% rise"

    @pytest.mark.asyncio
    async def test_a_quote_drifting_on_a_non_breaking_space_still_matches(
        self, docx_path: Path
    ):
        edit = _edit("Energy Supply chapter", "Energy Demand chapter", 1)

        plan, _ = await _plan(docx_path, [edit])
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        # The search string handed to Word keeps the document's own spacing.
        assert plan.planned[0].find == "Energy\u00a0Supply chapter"
        visible, _, _ = _read(docx_path)
        assert "The Energy Demand chapter reports" in visible

    @pytest.mark.asyncio
    async def test_an_empty_replacement_is_a_tracked_deletion(self, docx_path: Path):
        edit = _edit("for 2019", "", 1)

        plan, _ = await _plan(docx_path, [edit])
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        visible, original, revisions = _read(docx_path)
        assert "for 2019" not in visible
        assert "for 2019" in original
        assert [(r.type, r.text) for r in revisions] == [("deletion", "for 2019")]

    @pytest.mark.asyncio
    async def test_several_edits_in_one_paragraph_are_all_written(
        self, docx_path: Path
    ):
        first = _edit("Energy Supply", "Energy Demand", 1)
        second = _edit("14% rise", "18% rise", 1)

        plan, _ = await _plan(docx_path, [first, second])
        outcomes = await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        assert _statuses(outcomes) == {first.id: "applied", second.id: "applied"}
        visible, _, _ = _read(docx_path)
        assert "The Energy Demand chapter reports a 18% rise in output" in visible

    @pytest.mark.asyncio
    async def test_edits_in_different_paragraphs_are_all_written(self, docx_path: Path):
        first = _edit("14% rise", "18% rise", 1)
        second = _edit("Shared caption text", "Shared caption wording", 5)

        plan, _ = await _plan(docx_path, [first, second])
        outcomes = await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        assert _statuses(outcomes) == {first.id: "applied", second.id: "applied"}
        visible, _, _ = _read(docx_path)
        assert "18% rise" in visible
        assert "Shared caption wording" in visible


class TestParagraphResolution:
    @pytest.mark.asyncio
    async def test_a_repeated_phrase_uses_the_occurrence_the_source_settles(
        self, docx_path: Path
    ):
        # Unique in the markdown as `**Figure 3**`, the second of two identical
        # spans once the emphasis is stripped.
        edit = _edit("**Figure 3**", "Figure 4", 3)

        plan, _ = await _plan(docx_path, [edit])
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        assert plan.planned[0].occurrence == 1
        visible, _, _ = _read(docx_path)
        assert "Figure 3 and Figure 4 close the section." in visible

    @pytest.mark.asyncio
    async def test_a_table_cell_never_takes_a_body_paragraph_redline(
        self, docx_path: Path
    ):
        edit = _edit("Shared caption text", "Shared caption wording", 5)

        plan, _ = await _plan(docx_path, [edit])
        await apply_tracked_changes(
            str(docx_path), plan.planned, workspace_root=str(docx_path.parent)
        )

        reopened = PythonDocxDocument(str(docx_path))
        assert reopened.tables[0].cell(0, 0).text == "Shared caption text"
        assert reopened.tables[0].cell(0, 1).text == "Other cell"
        visible, _, _ = _read(docx_path)
        assert "Shared caption wording" in visible

    @pytest.mark.asyncio
    async def test_a_quote_the_paragraph_does_not_carry_is_not_found(
        self, docx_path: Path
    ):
        # On the markdown line, absent from the converted paragraph.
        edit = _edit("per Table 2", "per Table 3", 1)

        plan, _ = await _plan(docx_path, [edit])

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "not_found"}

    @pytest.mark.asyncio
    async def test_an_edit_whose_line_maps_to_no_paragraph_is_not_found(
        self, docx_path: Path
    ):
        edit = _edit("14% rise", "18% rise", 1)

        plan, _ = await _plan(docx_path, [edit], paragraph_line_ranges={})

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "not_found"}

    @pytest.mark.asyncio
    async def test_a_quote_the_line_carries_twice_never_reaches_the_plan(
        self, docx_path: Path
    ):
        edit = _edit("Figure 3", "Figure 4", 3)

        plan, _ = await _plan(docx_path, [edit])

        assert plan.planned == []
        assert _statuses(plan.outcomes) == {edit.id: "unlocatable"}


_REFERENCE_MARKDOWN = "Smith et al. 2019. Annual report."


@pytest.fixture
def reference_docx_path(tmp_path: Path) -> Path:
    document = PythonDocxDocument()
    document.add_paragraph(_REFERENCE_MARKDOWN)
    path = tmp_path / "reference.docx"
    document.save(str(path))
    return path


class TestAMidLineNumberIsNotAListMarker:
    @pytest.mark.asyncio
    async def test_a_quote_opening_with_a_year_keeps_it(
        self, reference_docx_path: Path
    ):
        edit = _edit("2019. Annual report", "2021. Annual report", 1)

        plan, _ = await _plan(
            reference_docx_path,
            [edit],
            paragraph_line_ranges={0: (1, 1)},
            markdown=_REFERENCE_MARKDOWN,
        )
        await apply_tracked_changes(
            str(reference_docx_path),
            plan.planned,
            workspace_root=str(reference_docx_path.parent),
        )

        assert _statuses(plan.outcomes) == {edit.id: "applied"}
        assert plan.planned[0].find == "2019. Annual report"
        visible, original, _ = _read(reference_docx_path)
        assert "Smith et al. 2021. Annual report." in visible
        assert "Smith et al. 2019. Annual report." in original
