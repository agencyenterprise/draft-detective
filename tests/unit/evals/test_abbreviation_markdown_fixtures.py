import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from inspect_ai.agent import AgentState
from inspect_ai.model import ChatMessageUser

from evals_inspectai.common import api_solver
from evals_inspectai.e2e.abbreviation_checker import generate_test_variants as generator
from evals_inspectai.e2e.abbreviation_checker.abbreviation_checker_e2e import (
    abbreviation_checker_entries,
    abbreviation_checker_from_files,
)
from evals_inspectai.e2e.abbreviation_checker.fixtures import document_files


def entry():
    return {
        "abbr": "AI",
        "inline_definition": "artificial intelligence",
        "occurrence_number": 1,
        "line_start": 3,
        "line_end": 3,
        "ignored": False,
        "ignored_reason": None,
        "abbreviations_section_definition": "artificial intelligence",
    }


def test_discovery_prefers_markdown_without_losing_explicit_docx_selection(tmp_path):
    for name in (
        "one.md",
        "one.markdown",
        "one.docx",
        "two.docx",
        "~$temp.docx",
        "references.json",
        "README.md",
    ):
        (tmp_path / name).touch()
    assert [p.name for p in document_files(tmp_path)] == ["one.md", "two.docx"]
    assert document_files(tmp_path, "one.docx") == [tmp_path / "one.docx"]


def test_public_example_has_matching_line_anchors_and_two_abbreviation_samples():
    directory = (
        Path(__file__).resolve().parents[3]
        / "evals_inspectai/e2e/abbreviation_checker/data/example"
    )
    task = abbreviation_checker_entries(data_dir=str(directory))
    assert len(task.dataset) == 2
    text = (directory / "baseline.md").read_text().splitlines()
    for sample in task.dataset:
        entries = json.loads(sample.target)
        assert len(entries) == 3
        for item in entries:
            assert item["abbr"] in text[item["line_start"] - 1]
            if item["inline_definition"]:
                assert item["inline_definition"] in text[item["line_start"] - 1]


@pytest.mark.parametrize(
    "factory", [abbreviation_checker_entries, abbreviation_checker_from_files]
)
def test_partial_references_fail_instead_of_silently_skipping_documents(
    tmp_path, factory
):
    (tmp_path / "one.md").write_text("AI")
    (tmp_path / "references.json").write_text("{}")
    with pytest.raises(ValueError, match="Missing reference"):
        factory(data_dir=str(tmp_path))


@pytest.mark.parametrize(
    "factory", [abbreviation_checker_entries, abbreviation_checker_from_files]
)
def test_agent_tasks_select_markdown_reference_and_do_not_double_dataset(
    tmp_path, factory
):
    (tmp_path / "one.md").write_text("# Title\n\nartificial intelligence (AI)")
    (tmp_path / "one.docx").touch()
    reference = {"abbreviations": [entry()], "abbreviations_section_found": True}
    (tmp_path / "references.json").write_text(
        json.dumps({"one.md": reference, "one.docx": reference})
    )
    task = factory(data_dir=str(tmp_path))
    assert len(task.dataset) == 1
    assert task.dataset[0].input == str(tmp_path / "one.md")
    assert str(task.dataset[0].id).startswith("one.md")


@pytest.mark.asyncio
async def test_agent_uploads_markdown_verbatim(monkeypatch, tmp_path):
    text = "# Title\n\n| AI | artificial intelligence |\n"
    path = tmp_path / "one.md"
    path.write_text(text, encoding="utf-8")
    start = AsyncMock(return_value="project")
    monkeypatch.setattr(api_solver, "create_project_and_start_workflows", start)
    monkeypatch.setattr(
        api_solver, "poll_until_complete", AsyncMock(return_value={"state": {}})
    )
    monkeypatch.setattr(
        api_solver, "sample_active", lambda: SimpleNamespace(epoch=1, eval_id="eval-1")
    )
    agent = api_solver.api_workflow_agent_file_cached(
        "abbreviation_scan_v2", {}, api_base_url="https://example.test"
    )
    await agent(AgentState(messages=[ChatMessageUser(content=str(path))]))
    assert start.call_args.kwargs["file_content"] == text.encode("utf-8")
    assert start.call_args.kwargs["file_name"] == "one.md"


def test_markdown_generator_edits_only_target_line_and_reference(tmp_path):
    text = "# Title\n\nartificial intelligence (AI) and AI.\n\nartificial intelligence (AI) again.\n"
    source = tmp_path / "baseline.md"
    source.write_text(text)
    first = entry()
    later = {**first, "occurrence_number": 3, "line_start": 5, "line_end": 5}
    original = {"abbreviations_section_found": False, "abbreviations": [first, later]}
    references = tmp_path / "references.json"
    references.write_text(json.dumps({source.name: original}))
    assert generator.generate(source, references, tmp_path / "variants") == 21
    output = tmp_path / "variants" / "single"
    variant = (output / "AI_no_def.md").read_text()
    assert variant.splitlines()[2] == "AI and AI."
    assert variant.splitlines()[4] == text.splitlines()[4]
    assert len(variant.splitlines()) == len(text.splitlines())
    reference = json.loads((output / "references.json").read_text())["AI_no_def.md"]
    assert reference["abbreviations"][0]["inline_definition"] == ""
    assert reference["abbreviations"][1] == later
    assert reference["abbreviations_section_found"] is False


def test_bad_generator_reference_fails_before_writing(tmp_path):
    source = tmp_path / "baseline.md"
    source.write_text("# Title\n\nAI without definition\n")
    references = tmp_path / "references.json"
    references.write_text(json.dumps({source.name: {"abbreviations": [entry()]}}))
    with pytest.raises(ValueError, match="Expected exactly one definition"):
        generator.generate(source, references, tmp_path / "variants")
    assert not (tmp_path / "variants").exists()


@pytest.mark.asyncio
async def test_migration_preserves_conversion_and_references_and_refuses_edits(
    tmp_path,
):
    from lib.services.converters.base import convert_to_markdown

    doc = Document()
    doc.add_heading("Title", level=1)
    doc.add_paragraph("artificial intelligence (AI)")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "AI"
    table.cell(0, 1).text = "artificial intelligence"
    path = tmp_path / "baseline.docx"
    doc.save(path)
    original = {"abbreviations": [entry()], "abbreviations_section_found": True}
    references = tmp_path / "references.json"
    references.write_text(json.dumps({path.name: original}))
    assert await generator.migrate_docx(tmp_path) == 1
    markdown = path.with_suffix(".md")
    assert markdown.read_text() == await convert_to_markdown(str(path))
    assert "|" in markdown.read_text()
    refs = json.loads(references.read_text())
    assert refs[path.name] == refs[markdown.name] == original
    assert path.exists()
    assert await generator.migrate_docx(tmp_path) == 1  # Idempotent snapshot.
    markdown.write_text("Manually edited fixture")
    with pytest.raises(ValueError, match="Refusing to overwrite"):
        await generator.migrate_docx(tmp_path)
    assert markdown.read_text() == "Manually edited fixture"
