"""The chat endpoints, with the model, checkpointer and database stubbed out.

What is pinned: who gets past the door, what the picker and slash-command menu
are fed, that a turn streams as server-sent events built from the agent's
events, that history and files come from the checkpointer, that deleting a
thread deletes its checkpoints, that a title lands on the thread, and how
attachment extraction fails.
"""

import io
import uuid
from datetime import datetime, timezone
from typing import Any, AsyncIterator
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage, HumanMessage

from lib.agents import chat_agent as chat_agent_module
from lib.api.auth import get_current_user
from lib.api.routers import chat
from lib.models.chat_thread import ChatThread
from lib.models.user import User, UserRole

USER = User(
    id=uuid.uuid4(),
    email="reviewer@example.com",
    name="Reviewer",
    role=UserRole.USER,
    show_experimental_features=False,
)
THREAD_ID = uuid.uuid4()


def _thread(title: str | None = None) -> ChatThread:
    now = datetime.now(timezone.utc)
    return ChatThread(
        id=THREAD_ID, user_id=USER.id, title=title, created_at=now, last_updated_at=now
    )


def _client(authenticated: bool = True) -> TestClient:
    app = FastAPI()
    app.include_router(chat.router)
    if authenticated:
        app.dependency_overrides[get_current_user] = lambda: USER
    return TestClient(app, raise_server_exceptions=False)


def _frames(body: str) -> list[str]:
    return [frame for frame in body.split("\n\n") if frame]


def _turn(events: list[dict[str, Any]], calls: list[dict[str, Any]]):
    """A stand-in for ``run_chat_turn``: records its arguments, replays events."""

    async def run_chat_turn(**kwargs: Any) -> AsyncIterator[dict[str, Any]]:
        calls.append(kwargs)
        for event in events:
            yield event

    return run_chat_turn


class TestGuard:
    def test_anonymous_requests_are_refused(self) -> None:
        client = _client(authenticated=False)
        assert client.get("/api/chat/models").status_code == 401
        assert client.get("/api/chat/skills").status_code == 401
        assert client.get(f"/api/chat/threads/{THREAD_ID}/messages").status_code == 401
        assert (
            client.post(f"/api/chat/threads/{THREAD_ID}/stream", json={"message": "hi"}).status_code
            == 401
        )


class TestCatalogues:
    def test_models_come_from_the_allowlist_with_one_default(self) -> None:
        rows = _client().get("/api/chat/models").json()
        assert [row["id"] for row in rows] == [o.id for o in chat_agent_module.CHAT_MODELS]
        assert [row["is_default"] for row in rows].count(True) == 1
        assert rows[0]["id"] == chat_agent_module.DEFAULT_CHAT_MODEL.name

    def test_skills_leave_out_the_excluded_ones(self) -> None:
        names = {row["name"] for row in _client().get("/api/chat/skills").json()}
        assert "reference-validation" in names
        assert "voice-and-tone" in names
        assert names.isdisjoint(chat_agent_module.EXCLUDED_CHAT_SKILLS)


class TestStream:
    EVENTS: list[dict[str, Any]] = [
        {"t": "reasoning", "v": "hm"},
        {"t": "tool", "id": "c1", "name": "read_file", "args": {"file_path": "/attachments/a.md"}},
        {"t": "tool_result", "id": "c1", "result": "text"},
        {"t": "text", "v": "Done."},
    ]

    def test_a_turn_streams_the_agent_events(self) -> None:
        calls: list[dict[str, Any]] = []
        with (
            patch.object(chat.chat_thread_service, "touch_thread", AsyncMock(return_value=_thread())) as touch,
            patch.object(chat, "run_chat_turn", _turn(self.EVENTS, calls)),
        ):
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/stream",
                json={
                    "model": "gpt-5.6-luna",
                    "message": "Check the methods.",
                    "message_id": "page-id-1",
                    "attachments": [{"name": "a.pdf", "text": "body"}],
                },
            )

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/event-stream")
        assert [frame.split("\n")[0][:30] for frame in _frames(response.text)] == [
            'data: {"t": "reasoning", "v": ',
            'data: {"t": "tool", "id": "c1"',
            'data: {"t": "tool_result", "id',
            'data: {"t": "text", "v": "Done',
        ]

        touch.assert_awaited_once()
        (call,) = calls
        assert call["thread_id"] == str(THREAD_ID)
        assert call["user_id"] == str(USER.id)
        assert call["model"].name == "gpt-5.6-luna"
        assert call["api_key"] is None
        assert call["text"] == "Check the methods."
        assert call["message_id"] == "page-id-1"
        assert [(a.name, a.text) for a in call["attachments"]] == [("a.pdf", "body")]

    def test_an_unknown_model_falls_back_to_the_default(self) -> None:
        calls: list[dict[str, Any]] = []
        with (
            patch.object(chat.chat_thread_service, "touch_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "run_chat_turn", _turn([], calls)),
        ):
            _client().post(
                f"/api/chat/threads/{THREAD_ID}/stream",
                json={"model": "gpt-4.1", "message": "hi"},
            )
        assert calls[0]["model"] == chat_agent_module.DEFAULT_CHAT_MODEL

    def test_an_empty_turn_is_rejected_before_touching_anything(self) -> None:
        with patch.object(chat.chat_thread_service, "touch_thread", AsyncMock()) as touch:
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/stream", json={"message": "   "}
            )
        assert response.status_code == 422
        touch.assert_not_awaited()

    def test_an_attachment_alone_is_a_valid_turn(self) -> None:
        calls: list[dict[str, Any]] = []
        with (
            patch.object(chat.chat_thread_service, "touch_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "run_chat_turn", _turn([], calls)),
        ):
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/stream",
                json={"attachments": [{"name": "a.pdf", "text": "body"}]},
            )
        assert response.status_code == 200 and calls[0]["text"] == ""

    def test_a_failure_mid_turn_ends_the_stream_with_an_error_event(self) -> None:
        async def failing(**kwargs: Any) -> AsyncIterator[dict[str, Any]]:
            yield {"t": "text", "v": "part"}
            raise RuntimeError("rate limited")

        with (
            patch.object(chat.chat_thread_service, "touch_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "run_chat_turn", failing),
        ):
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/stream", json={"message": "hi"}
            )
        assert response.status_code == 200
        assert _frames(response.text)[-1] == 'data: {"t": "error", "v": "rate limited"}'


class TestHistory:
    def test_messages_come_from_the_checkpointer_in_ui_shape(self) -> None:
        stored = [
            HumanMessage(content="hi\n\n[Attached ...]", id="h1", additional_kwargs={"user_text": "hi", "attachments": []}),
            AIMessage(content=[{"type": "text", "text": "hello"}], id="a1"),
        ]
        with (
            patch.object(chat.chat_thread_service, "get_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "load_thread_messages", AsyncMock(return_value=stored)) as load,
        ):
            response = _client().get(f"/api/chat/threads/{THREAD_ID}/messages")

        assert response.status_code == 200
        assert response.json() == [
            {"id": "h1", "type": "human", "content": "hi", "additional_kwargs": {"attachments": []}},
            {"id": "a1", "type": "ai", "content": [{"type": "text", "text": "hello"}], "tool_calls": []},
        ]
        load.assert_awaited_once_with(str(THREAD_ID))

    def test_only_attachments_are_served(self) -> None:
        """The same filesystem holds the mounted skills; those are not the page's to read."""
        with (
            patch.object(chat.chat_thread_service, "get_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "read_thread_file", AsyncMock(return_value="secret")) as read,
        ):
            for path in ("/skills/issues/SKILL.md", "/report.md", "attachments/a.md", "/attachmentsX/a.md"):
                assert _client().get(f"/api/chat/threads/{THREAD_ID}/files", params={"path": path}).status_code == 404
        read.assert_not_awaited()

    def test_a_thread_file_is_served_or_404(self) -> None:
        with (
            patch.object(chat.chat_thread_service, "get_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "read_thread_file", AsyncMock(side_effect=["one\ntwo", None])),
        ):
            found = _client().get(f"/api/chat/threads/{THREAD_ID}/files", params={"path": "/attachments/a.md"})
            missing = _client().get(f"/api/chat/threads/{THREAD_ID}/files", params={"path": "/nope"})
        assert found.status_code == 200
        assert found.json() == {"path": "/attachments/a.md", "content": "one\ntwo"}
        assert missing.status_code == 404

    def test_deleting_a_thread_deletes_its_checkpoints(self) -> None:
        with (
            patch.object(chat.chat_thread_service, "delete_thread", AsyncMock()) as delete_row,
            patch.object(chat, "delete_thread_state", AsyncMock()) as delete_state,
        ):
            response = _client().delete(f"/api/chat/threads/{THREAD_ID}")
        assert response.status_code == 204
        delete_row.assert_awaited_once()
        delete_state.assert_awaited_once_with(str(THREAD_ID))


class TestTitle:
    def test_the_generated_title_is_stored_on_the_thread(self) -> None:
        with (
            patch.object(chat.chat_thread_service, "get_thread", AsyncMock(return_value=_thread())),
            patch.object(chat, "generate_title", AsyncMock(return_value="Citation check")),
            patch.object(
                chat.chat_thread_service,
                "rename_thread",
                AsyncMock(return_value=_thread(title="Citation check")),
            ) as rename,
        ):
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/title",
                json={"messages": [{"role": "user", "content": "check my citations"}]},
            )
        assert response.status_code == 200
        assert response.json()["title"] == "Citation check"
        assert rename.call_args.kwargs["title"] == "Citation check"


    def test_a_titled_thread_keeps_its_title(self) -> None:
        """assistant-ui can ask twice around the first run; the second ask is a no-op."""
        with (
            patch.object(chat.chat_thread_service, "get_thread", AsyncMock(return_value=_thread(title="Kept"))),
            patch.object(chat, "generate_title", AsyncMock()) as generate,
            patch.object(chat.chat_thread_service, "rename_thread", AsyncMock()) as rename,
        ):
            response = _client().post(
                f"/api/chat/threads/{THREAD_ID}/title",
                json={"messages": [{"role": "user", "content": "hi"}]},
            )
        assert response.status_code == 200 and response.json()["title"] == "Kept"
        generate.assert_not_awaited()
        rename.assert_not_awaited()


class TestExtract:
    def test_a_docx_comes_back_as_markdown(self) -> None:
        document = Document()
        document.add_heading("Findings", level=1)
        document.add_paragraph("The effect was small.")
        buffer = io.BytesIO()
        document.save(buffer)

        response = _client().post(
            "/api/chat/extract", files={"file": ("draft.docx", buffer.getvalue())}
        )
        assert response.status_code == 200
        assert "Findings" in response.json()["text"]
        assert "The effect was small." in response.json()["text"]

    def test_other_types_are_refused(self) -> None:
        response = _client().post("/api/chat/extract", files={"file": ("notes.txt", b"hello")})
        assert response.status_code == 415

    def test_an_empty_document_is_reported(self) -> None:
        buffer = io.BytesIO()
        Document().save(buffer)
        response = _client().post(
            "/api/chat/extract", files={"file": ("blank.docx", buffer.getvalue())}
        )
        assert response.status_code == 422
